from .base import TextParser, ParsedQuestion


class WordParser:
    def parse_file(self, filepath: str) -> list[ParsedQuestion]:
        from docx import Document

        doc = Document(filepath)
        lines = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                lines.append(text)
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))

        text = "\n".join(lines)
        return TextParser().parse(text)

    def extract_bank_name(self, filepath: str) -> str:
        from docx import Document

        doc = Document(filepath)
        text = "\n".join(p.text for p in doc.paragraphs)
        return TextParser().extract_bank_name(text)
